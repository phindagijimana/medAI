"""Human review of LEO output: record decisions, then merge them into the result table.

A reviewer looks at one value at a time (one phenotype or question answer in one note) and decides:

    accept    the value is right                      final value = LEO's value
    correct   the value is wrong, here is the right   final value = what the reviewer typed
    reject    the value is wrong or unsupported       final value = blank  (Deny in the Validator)
    add       LEO missed a value the note states      final value = what the reviewer typed
              (only for ``not_found``; put the supporting words in ``comment``)

Optional tags (stored on the same files; extract does not read them): ``why_tag``, ``reader_preferred``
(rules/model/both/neither on a disagreement), ``quote_ok``. Deny and Correct need ``why_tag``.

Decisions are kept in two files next to the result (both contain note text, so treat them like the input):

    <out>.reviews.jsonl     append-only log: every decision ever made, with LEO's original proposal
    <out>.reviewed.csv      the latest decision per value, one row each (REVIEWED_COLUMNS); this is the
                            file you can open in Excel, check, or edit by hand

LEO's own columns are never changed.  ``merge`` writes a copy of the result table with two columns added
after each reviewed field:

    <field>_final    the value to use: the reviewer's, or LEO's if nobody objected (see below)
    <field>_review   where that value came from:
                     accepted_by_reviewer | corrected | rejected | added_by_reviewer |
                     auto_accepted | unreviewed | (blank)

``auto_accepted`` means LEO accepted it by its own checks and nobody looked at it (not a human decision).
``unreviewed`` values that need review are left blank in ``_final`` by default, because LEO did not trust
them; ``unreviewed="proposal"`` keeps LEO's value there instead (the ``_review`` column still says
"unreviewed").

A missed value (``not_found``) is offered under scope ``misses`` (Review tab: “Values LEO missed”,
``leo review --misses``).  Review files that include adds can be used for recall on the notes that were
reviewed; they still say nothing about notes nobody looked at.
"""

from __future__ import annotations

import hashlib
import json
import threading
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Optional

import pandas as pd

from services.leo_tabular import _safe

DECISIONS = ("accept", "correct", "reject", "add")
FLAGGED = ("needs_review", "rejected")                  # statuses a reviewer should look at
NOT_REVIEWABLE = ("", "not_found", "not_documented", "error")
WHY_TAGS = ("wrong_value", "wrong_status", "quote_does_not_support", "wrong_reader",
            "copy_forward", "missed", "other")
READER_PREFERRED = ("rules", "model", "both", "neither")
SCOPES = ("flagged", "misses", "all", "accepted_audit")
REVIEWED_COLUMNS = ["row", "doc_id", "field", "label", "proposed_value", "proposed_status", "proposed_why",
                    "proposed_evidence", "proposed_source", "decision", "final_value", "reviewer",
                    "reviewed_utc", "comment", "why_tag", "reader_preferred", "quote_ok", "leo_version"]
MAX_VALUE = 500
MAX_COMMENT = 500
_LOCK = threading.Lock()


@dataclass(frozen=True)
class Field:
    id: str                     # logical id: a phenotype id or a question id
    label: str
    value: str                  # table column names (may differ from ``id`` if the input had a clash)
    status: str
    evidence: str
    why: str = ""
    source: str = ""


def fields_from(column_map: dict, phenotypes: Iterable[Any], asks: Iterable[Any]) -> list[Field]:
    g = lambda c: column_map.get(c, c)
    out = [Field(p.id, p.label, g(p.id), g(f"{p.id}_status"), g(f"{p.id}_evidence"), g(f"{p.id}_why"),
                 g(f"{p.id}_source")) for p in phenotypes]
    out += [Field(a.id, a.question, g(a.id), g(f"{a.id}_status"), g(f"{a.id}_evidence")) for a in asks]
    return out


def infer_fields(table: pd.DataFrame) -> list[Field]:
    """Fields of a LEO table from its column names (when the run's provenance is not at hand)."""
    cols = set(table.columns)
    out = []
    for c in table.columns:
        if c.startswith("leo_") or f"{c}_status" not in cols or f"{c}_evidence" not in cols:
            continue
        out.append(Field(c, c, c, f"{c}_status", f"{c}_evidence",
                         f"{c}_why" if f"{c}_why" in cols else "", f"{c}_source" if f"{c}_source" in cols else ""))
    return out


def doc_column(table: pd.DataFrame) -> str:
    return next((c for c in table.columns if c.startswith("leo_doc_id")), "")


def _cell(table: pd.DataFrame, i: int, col: str) -> str:
    return "" if not col or col not in table.columns else str(table.iat[i, table.columns.get_loc(col)])


def reviewable(status: str, scope: str) -> bool:
    if scope == "flagged":
        return status in FLAGGED
    if scope == "misses":
        return status == "not_found"
    if scope == "all":
        return status not in NOT_REVIEWABLE
    if scope == "accepted_audit":
        return status == "accepted"
    raise ValueError("scope must be flagged, misses, all or accepted_audit")


def in_accepted_audit(doc_id: str, field: str, row: int, seed: str = "") -> bool:
    """Deterministic ~10% sample of auto-accepted cells (audit of silent errors)."""
    raw = f"{seed}|{doc_id}|{row}|{field}".encode()
    return int(hashlib.md5(raw).hexdigest()[:8], 16) % 10 == 0


def _column_reader(table: pd.DataFrame):
    """cell(i, col) backed by per-column string lists, built on first use (fast for whole-table scans)."""
    cache: dict = {}

    def cell(i: int, col: str) -> str:
        if not col or col not in table.columns:
            return ""
        if col not in cache:
            cache[col] = table[col].astype(str).tolist()
        return cache[col][i]
    return cell


def proposal(table: pd.DataFrame, f: Field, row: int, cell=None) -> dict:
    """What LEO proposed for field ``f`` in 1-based ``row``."""
    cell = cell or (lambda i, col: _cell(table, i, col))
    i, dc = row - 1, doc_column(table)
    return {"row": row, "doc_id": cell(i, dc) or f"row{row}", "field": f.id, "label": f.label,
            "value": cell(i, f.value), "status": cell(i, f.status),
            "why": cell(i, f.why), "evidence": cell(i, f.evidence),
            "source": cell(i, f.source),
            "alternative": cell(i, f"{f.id}_alternative"),
            "code": cell(i, f"{f.id}_code"),
            "code_system": cell(i, f"{f.id}_system"),
            "code_status": cell(i, f"{f.id}_code_status")}


def items(table: pd.DataFrame, fields: list[Field], latest: dict, scope: str = "flagged",
          audit_seed: str = "") -> list[dict]:
    """Values to review (in table order), each with the current decision if there is one."""
    if scope not in SCOPES:
        raise ValueError("scope must be flagged, misses, all or accepted_audit")
    out = []
    cell = _column_reader(table)
    lcol = next((c for c in table.columns if c.startswith("leo_status")), "")
    for i in range(len(table)):
        if lcol and cell(i, lcol).startswith("error"):
            continue
        for f in fields:
            if f.status not in table.columns or not reviewable(cell(i, f.status), scope):
                continue
            it = proposal(table, f, i + 1, cell)
            if scope == "accepted_audit" and not in_accepted_audit(it["doc_id"], f.id, i + 1, audit_seed):
                continue
            ev = latest.get((i + 1, f.id))
            it.update({"decision": ev["decision"] if ev else "", "final": ev["final_value"] if ev else "",
                       "comment": ev["comment"] if ev else "",
                       "why_tag": ev.get("why_tag", "") if ev else "",
                       "reader_preferred": ev.get("reader_preferred", "") if ev else ""})
            out.append(it)
    return out


def counts(its: list[dict]) -> dict:
    c = {"total": len(its), "reviewed": 0, "accept": 0, "correct": 0, "reject": 0, "add": 0}
    for it in its:
        if it["decision"]:
            c["reviewed"] += 1
            if it["decision"] in c:
                c[it["decision"]] += 1
    return c


def make_event(prop: dict, decision: str, value: str, reviewer: str, comment: str, version: str,
               why_tag: str = "", reader_preferred: str = "") -> dict:
    """A decision record.  ``decision='undo'`` withdraws an earlier decision."""
    reviewer = (reviewer or "").strip()
    if not reviewer:
        raise ValueError("enter a reviewer name or ID (a pseudonym or initials; not a patient name)")
    if len(reviewer) > 60:
        raise ValueError("reviewer is too long")
    decision = (decision or "").strip().lower()
    if decision not in DECISIONS + ("undo",):
        raise ValueError(f"decision must be one of {', '.join(DECISIONS)}")
    value, comment = (value or "").strip(), (comment or "").strip()
    why_tag, reader_preferred = (why_tag or "").strip().lower(), (reader_preferred or "").strip().lower()
    if len(value) > MAX_VALUE or len(comment) > MAX_COMMENT:
        raise ValueError("value or comment is too long")
    if decision == "correct":
        if not value:
            raise ValueError("a correction needs the right value; use reject if there is none")
        if value == prop["value"]:
            decision = "accept"
    if decision == "add":
        if not value:
            raise ValueError("adding a missed value needs the value")
        if (prop.get("status") or "") != "not_found":
            raise ValueError("add is for a missed value; use correct to change one LEO found")
        if not why_tag:
            why_tag = "missed"
    if decision != "undo":
        if why_tag and why_tag not in WHY_TAGS:
            raise ValueError(f"why must be one of {', '.join(WHY_TAGS)}")
        if decision in ("correct", "reject") and not why_tag:
            raise ValueError("choose why before Deny or Correct (wrong_value, wrong_status, "
                             "quote_does_not_support, wrong_reader, copy_forward, missed, other)")
        if reader_preferred and reader_preferred not in READER_PREFERRED:
            raise ValueError(f"reader_preferred must be one of {', '.join(READER_PREFERRED)}")
        if (prop.get("alternative") or "") and decision in ("accept", "correct", "reject") and not reader_preferred:
            raise ValueError("this cell is a disagreement; say which reader was right (rules, model, both, neither)")
    if why_tag == "quote_does_not_support":
        quote_ok = "no"
    elif decision == "accept":
        quote_ok = "yes"
    else:
        quote_ok = ""
    final = {"accept": prop["value"], "correct": value, "reject": "", "undo": "", "add": value}[decision]
    return {"row": prop["row"], "doc_id": prop["doc_id"], "field": prop["field"], "label": prop["label"],
            "proposed_value": prop["value"], "proposed_status": prop["status"], "proposed_why": prop["why"],
            "proposed_evidence": prop["evidence"], "proposed_source": prop["source"], "decision": decision,
            "final_value": final, "reviewer": reviewer,
            "reviewed_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "comment": comment, "why_tag": why_tag, "reader_preferred": reader_preferred,
            "quote_ok": quote_ok, "leo_version": version}


class ReviewStore:
    """Append-only decision log plus the readable ``reviewed.csv`` (latest decision per value)."""

    def __init__(self, log_path: str | Path, csv_path: str | Path):
        self.log, self.csv = Path(log_path), Path(csv_path)

    def latest(self) -> dict:
        out: dict = {}
        if not self.log.exists():
            return out
        for line in self.log.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            ev = json.loads(line)
            key = (int(ev["row"]), ev["field"])
            if ev["decision"] == "undo":
                out.pop(key, None)
            else:
                out[key] = ev
        return out

    def add(self, events: list[dict]) -> dict:
        with _LOCK:
            self.log.parent.mkdir(parents=True, exist_ok=True)
            with self.log.open("a", encoding="utf-8") as fh:
                for ev in events:
                    fh.write(json.dumps(ev, ensure_ascii=False) + "\n")
            latest = self.latest()
            self._write_csv(latest)
            return latest

    def _write_csv(self, latest: dict) -> None:
        rows = [{c: (_safe(ev.get(c, "")) if c in ("proposed_value", "final_value", "comment") else ev.get(c, ""))
                 for c in REVIEWED_COLUMNS} for _, ev in sorted(latest.items(), key=lambda kv: (kv[0][0], kv[0][1]))]
        pd.DataFrame(rows, columns=REVIEWED_COLUMNS).to_csv(self.csv, index=False, encoding="utf-8-sig")


def read_reviewed(path: str | Path, table: Optional[pd.DataFrame] = None) -> dict:
    """Load a reviewed.csv (possibly edited by hand) as {(row, field): record}; check it fits ``table``."""
    p = Path(path)
    df = pd.read_csv(p, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    missing = [c for c in ("row", "field", "decision") if c not in df.columns]
    if missing:
        raise ValueError(f"{p.name}: missing column(s) {missing}; expected {REVIEWED_COLUMNS}")
    out: dict = {}
    dc = doc_column(table) if table is not None else ""
    for n, r in df.iterrows():
        dec = str(r["decision"]).strip().lower()
        if not dec:
            continue                                         # a row not yet decided
        if dec not in DECISIONS:
            raise ValueError(f"{p.name} line {n + 2}: decision {dec!r} is not one of {', '.join(DECISIONS)}")
        try:
            row = int(r["row"])
        except ValueError:
            raise ValueError(f"{p.name} line {n + 2}: row {r['row']!r} is not a number")
        if table is not None:
            if not 1 <= row <= len(table):
                raise ValueError(f"{p.name} line {n + 2}: row {row} is outside the table ({len(table)} rows); "
                                 "was this file made for a different result?")
            doc = str(r.get("doc_id", "")).strip()
            if doc and dc and doc != _cell(table, row - 1, dc):
                raise ValueError(f"{p.name} line {n + 2}: doc_id {doc!r} does not match the table's "
                                 f"{_cell(table, row - 1, dc)!r} at row {row}; was this file made for a different result?")
        final = str(r.get("final_value", "")).strip()
        if dec in ("correct", "add") and not final:
            raise ValueError(f"{p.name} line {n + 2}: a {'correction' if dec == 'correct' else 'missed value'} needs final_value")
        out[(row, str(r["field"]).strip())] = {"decision": dec, "final_value": final,
                                                "reviewer": str(r.get("reviewer", "")), "comment": str(r.get("comment", "")),
                                                "why_tag": str(r.get("why_tag", "")).strip(),
                                                "reader_preferred": str(r.get("reader_preferred", "")).strip(),
                                                "quote_ok": str(r.get("quote_ok", "")).strip()}
    return out


def merge(table: pd.DataFrame, fields: list[Field], latest: dict, unreviewed: str = "blank") -> pd.DataFrame:
    """Copy of ``table`` with ``<field>_final`` and ``<field>_review`` added after each reviewed field."""
    if unreviewed not in ("blank", "proposal"):
        raise ValueError("unreviewed must be blank or proposal")
    known = {f.id for f in fields}
    stray = sorted({fid for (_, fid) in latest if fid not in known})
    if stray:
        raise ValueError(f"reviewed file names fields that are not in the table: {stray}")
    existing = set(table.columns)

    def unique(name: str) -> str:
        while name in existing:
            name += "_leo"
        existing.add(name)
        return name

    after: dict[str, list[tuple[str, list[str]]]] = {}
    for f in fields:
        if f.value not in table.columns or f.status not in table.columns:
            continue
        finals, revs = [], []
        for i in range(len(table)):
            st, val = _cell(table, i, f.status), _cell(table, i, f.value)
            ev = latest.get((i + 1, f.id))
            if ev:
                d = ev["decision"]
                finals.append(val if d == "accept" and not ev["final_value"] else _safe(ev["final_value"]))
                revs.append({"accept": "accepted_by_reviewer", "correct": "corrected",
                             "reject": "rejected", "add": "added_by_reviewer"}[d])
            elif st == "accepted":
                finals.append(val); revs.append("auto_accepted")
            elif st in NOT_REVIEWABLE:
                finals.append(""); revs.append("")
            else:
                finals.append(val if unreviewed == "proposal" else ""); revs.append("unreviewed")
        anchor = f.why if f.why in table.columns else (f.evidence if f.evidence in table.columns else f.status)
        after.setdefault(anchor, []).extend([(unique(f"{f.id}_final"), finals), (unique(f"{f.id}_review"), revs)])
    cols: dict[str, Any] = {}
    for c in table.columns:
        cols[c] = table[c].tolist()
        for name, vals in after.get(c, []):
            cols[name] = vals
    return pd.DataFrame(cols)


def reviewed_frame(latest: dict) -> pd.DataFrame:
    rows = [{c: ev.get(c, "") for c in REVIEWED_COLUMNS} for _, ev in sorted(latest.items())]
    return pd.DataFrame(rows, columns=REVIEWED_COLUMNS)


def template(table: pd.DataFrame, fields: list[Field], scope: str = "flagged",
             audit_seed: str = "") -> pd.DataFrame:
    """A reviewed.csv to fill in by hand: one row per value to review, decision left blank.
    Fill ``decision`` (accept / correct / reject / add) and, for correct or add, ``final_value``; then ``merge_files``."""
    rows = []
    for it in items(table, fields, {}, scope, audit_seed=audit_seed):
        rows.append({"row": it["row"], "doc_id": it["doc_id"], "field": it["field"], "label": it["label"],
                     "proposed_value": it["value"], "proposed_status": it["status"], "proposed_why": it["why"],
                     "proposed_evidence": it["evidence"], "proposed_source": it["source"], "decision": "",
                     "final_value": "", "reviewer": "", "reviewed_utc": "", "comment": "",
                     "why_tag": "", "reader_preferred": "", "quote_ok": "", "leo_version": ""})
    return pd.DataFrame(rows, columns=REVIEWED_COLUMNS)


def load_table(path: str | Path) -> pd.DataFrame:
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        return pd.read_excel(p, sheet_name="LEO", dtype=str).fillna("")
    return pd.read_csv(p, dtype=str, keep_default_na=False, encoding="utf-8-sig")


def write_merged(merged: pd.DataFrame, reviews: pd.DataFrame, out: str | Path, *, source_xlsx: Optional[str | Path] = None) -> Path:
    """Write the merged table.  .csv: one file.  .xlsx: the source workbook with its LEO sheet replaced and a
    Reviews sheet added (Claims and About are kept)."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.suffix.lower() != ".xlsx":
        merged.to_csv(out, index=False, encoding="utf-8-sig")
        return out
    from openpyxl import Workbook, load_workbook
    wb = load_workbook(source_xlsx) if source_xlsx else Workbook()
    idx = wb.sheetnames.index("LEO") if "LEO" in wb.sheetnames else 0
    for name in ("LEO", "Reviews"):
        if name in wb.sheetnames:
            del wb[name]
    if not wb.sheetnames:
        wb.create_sheet("tmp")
    for name, df, pos in (("LEO", merged, idx), ("Reviews", reviews, None)):
        ws = wb.create_sheet(name, pos)
        ws.append(list(df.columns))
        for rec in df.itertuples(index=False):
            ws.append([_safe(v) if isinstance(v, str) else v for v in rec])
        ws.freeze_panes = "A2"
    if "tmp" in wb.sheetnames:
        del wb["tmp"]
    wb.save(out)
    return out


def merge_files(table_path: str | Path, reviewed_path: str | Path, out: Optional[str | Path] = None,
                unreviewed: str = "blank") -> Path:
    """Merge a reviewed.csv into a LEO result file (csv or xlsx); returns the path written."""
    tp = Path(table_path)
    if tp.suffix.lower() not in (".csv", ".xlsx"):
        raise ValueError("the result table must be a .csv or .xlsx file written by LEO")
    table = load_table(tp)
    fields = infer_fields(table)
    if not fields:
        raise ValueError(f"{tp.name} does not look like a LEO result (no <field>_status columns)")
    latest = read_reviewed(reviewed_path, table)
    merged = merge(table, fields, latest, unreviewed)
    out = Path(out) if out else tp.with_name(tp.stem + ".merged" + tp.suffix)
    return write_merged(merged, reviewed_frame(latest), out, source_xlsx=tp if tp.suffix.lower() == ".xlsx" else None)


def summarize_reviews(path: str | Path) -> str:
    """Per-field Accept/Correct/Deny rates and disagreement reader votes. Labels only; does not change extract."""
    p = Path(path)
    df = pd.read_csv(p, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if "decision" not in df.columns or "field" not in df.columns:
        raise ValueError(f"{p.name} is not a reviewed.csv")
    df = df[df["decision"].str.strip().str.lower().isin(DECISIONS)].copy()
    if df.empty:
        return f"{p.name}: no decisions yet."
    lines = [f"# Validator summary ({p.name})", "",
             "Labels only. Does not change extract or D17. Few-shot and LoRA/BIO are paused.", ""]
    lines.append("| field | n | accept | correct | reject | add |")
    lines.append("|---|---:|---:|---:|---:|---:|")
    for field, g in df.groupby(df["field"].str.strip(), sort=True):
        d = g["decision"].str.strip().str.lower()
        lines.append(f"| `{field}` | {len(g)} | {int((d=='accept').sum())} | {int((d=='correct').sum())} | "
                     f"{int((d=='reject').sum())} | {int((d=='add').sum())} |")
    if "why_tag" in df.columns:
        tags = df["why_tag"].str.strip().str.lower()
        tags = tags[tags != ""]
        if len(tags):
            lines += ["", "Why tags:"]
            for t, n in tags.value_counts().items():
                lines.append(f"- `{t}`: {n}")
    if "reader_preferred" in df.columns:
        rp = df[df["reader_preferred"].str.strip() != ""]
        if len(rp):
            lines += ["", "Which reader was right (disagreements):"]
            lines.append("| field | rules | model | both | neither |")
            lines.append("|---|---:|---:|---:|---:|")
            for field, g in rp.groupby(rp["field"].str.strip(), sort=True):
                v = g["reader_preferred"].str.strip().str.lower()
                lines.append(f"| `{field}` | {int((v=='rules').sum())} | {int((v=='model').sum())} | "
                             f"{int((v=='both').sum())} | {int((v=='neither').sum())} |")
            lines.append("")
            lines.append("To score routing and thresholds: `leo learn FILE.reviewed.csv`. "
                         "A policy file is written only with `--emit-policy` after the ≥100/field CI gate. D17 stays until then.")
    if "quote_ok" in df.columns:
        q = df["quote_ok"].str.strip().str.lower()
        n_yes, n_no = int((q == "yes").sum()), int((q == "no").sum())
        if n_yes or n_no:
            lines += ["", f"Quote support: yes {n_yes}, no {n_no} (feeds Loop B / quote-entail, not live)."]
    return "\n".join(lines) + "\n"


__all__ = ["Field", "fields_from", "infer_fields", "items", "counts", "proposal", "make_event", "ReviewStore",
           "read_reviewed", "merge", "merge_files", "template", "write_merged", "reviewed_frame", "load_table",
           "summarize_reviews", "in_accepted_audit",
           "REVIEWED_COLUMNS", "DECISIONS", "FLAGGED", "WHY_TAGS", "READER_PREFERRED", "SCOPES"]
