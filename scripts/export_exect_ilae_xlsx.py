"""Export ExECT-V2's ILAE→UMLS/SNOMEDCT XLSX to CSV for offline study.

Source (local clone, gitignored):
  references/exect_v2/Other/ILAE proposed map to UMLS and SNOMEDCT.xlsx

Output:
  references/exect_v2_ilae_map.csv  (also gitignored — this is redistributed
  reference material derived from ExECT-V2; we keep it local)

The XLSX contains ILAE-2017 concepts mapped to SNOMED CT IDs and UMLS CUIs
plus (in some columns) annotator notes. We flatten every non-empty sheet
into a single CSV with a `sheet` column so downstream code can filter.

Phase 0.2 of the MedCAT+Rules build plan (docs/MedCAT.md).
Reference-only; do not redistribute.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = (
    REPO_ROOT
    / "references"
    / "exect_v2"
    / "Other"
    / "ILAE proposed map to UMLS and SNOMEDCT.xlsx"
)
CSV_PATH = REPO_ROOT / "references" / "exect_v2_ilae_map.csv"


def _row_is_empty(row: tuple) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row)


def export() -> tuple[int, int, list[str]]:
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"XLSX not found: {XLSX_PATH}")

    wb = load_workbook(XLSX_PATH, data_only=True, read_only=True)

    all_columns: list[str] = ["sheet"]
    rows_out: list[dict] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            continue
        # Normalise headers → strings, blank cells become "col_<n>"
        norm_header = []
        for i, cell in enumerate(header):
            if cell is None or (isinstance(cell, str) and not cell.strip()):
                norm_header.append(f"col_{i}")
            else:
                norm_header.append(str(cell).strip())
        for col in norm_header:
            if col not in all_columns:
                all_columns.append(col)

        for row in rows_iter:
            if _row_is_empty(row):
                continue
            record = {"sheet": sheet_name}
            for col_name, cell in zip(norm_header, row):
                if cell is None:
                    continue
                # Coerce all cells to str to avoid CSV type-mixing surprises
                record[col_name] = str(cell).strip() if isinstance(cell, str) else str(cell)
            rows_out.append(record)

    CSV_PATH.parent.mkdir(parents=True, exist_ok=True)
    with CSV_PATH.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=all_columns, extrasaction="ignore")
        writer.writeheader()
        for record in rows_out:
            writer.writerow(record)

    return len(rows_out), len(all_columns), wb.sheetnames


if __name__ == "__main__":
    n_rows, n_cols, sheets = export()
    print(f"Sheets processed : {sheets}")
    print(f"Rows written     : {n_rows}")
    print(f"Columns written  : {n_cols}")
    print(f"Output path      : {CSV_PATH}")
    sys.exit(0)
